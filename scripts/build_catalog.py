"""종합쇼핑몰 품명별 수집 결과(data/raw/shopmall/*.csv)를 정부권장정책 엑셀과
사업자등록번호로 조인해 우대기업 품목 카탈로그(site/data/)를 만든다.

- 조인 키는 cntrctCorpBizno <-> 엑셀 사업자번호 뿐이다. 기업명 문자열 매칭은 하지 않는다.
- 우대기업이 아닌 업체의 품목은 저장하지 않는다(품명별 전체 건수만 summary에 남긴다).
- 원시 CSV가 없으면 가공 데이터를 만들지 않고 중단한다.
"""
import collections
import csv
import datetime as dt
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / 'data' / 'raw' / 'shopmall'
EXCEL_PATH = ROOT / '2026년 6월 30일 기준 정부권장정책.xlsx'
OUTPUT_DIR = ROOT / 'site' / 'data'
CATALOG_PATH = OUTPUT_DIR / 'catalog.json'
SUMMARY_PATH = OUTPUT_DIR / 'summary.json'

CHECK_DATE = dt.date.today().isoformat()
MAX_FILE_SIZE_BYTES = 2 * 1024 * 1024  # 2MB


def normalize_bizno(v):
    """Normalize business number to 10-digit numeric string with zero-padding."""
    if v is None:
        return ''
    s = str(v).replace('-', '').strip()
    if not s or not s.isdigit():
        return ''
    return s.zfill(10)


def normalize_date(v):
    """Normalize int(20241023), str, or datetime to YYYY-MM-DD."""
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    if not s or s.lower() == 'none':
        return None
    if s.isdigit() and len(s) == 8:
        return f"{s[:4]}-{s[4:6]}-{s[6:]}"
    if len(s) >= 10 and s[4] == '-' and s[7] == '-':
        return s[:10]
    return None


def extract_region(address):
    """Extract standard metropolitan city / province (광역 시도) only."""
    if not address:
        return ''
    first = str(address).strip().split()[0] if str(address).strip().split() else ''
    region_map = {
        '서울': '서울특별시', '서울특별시': '서울특별시', '서울시': '서울특별시',
        '경기': '경기도', '경기도': '경기도',
        '인천': '인천광역시', '인천광역시': '인천광역시', '인천시': '인천광역시',
        '강원': '강원특별자치도', '강원도': '강원특별자치도', '강원특별자치도': '강원특별자치도',
        '충북': '충청북도', '충청북도': '충청북도',
        '충남': '충청남도', '충청남도': '충청남도',
        '대전': '대전광역시', '대전광역시': '대전광역시', '대전시': '대전광역시',
        '세종': '세종특별자치시', '세종특별자치시': '세종특별자치시', '세종시': '세종특별자치시',
        '전북': '전북특별자치도', '전라북도': '전북특별자치도', '전북특별자치도': '전북특별자치도',
        '전남': '전라남도', '전라남도': '전라남도',
        '광주': '광주광역시', '광주광역시': '광주광역시', '광주시': '광주광역시',
        '경북': '경상북도', '경상북도': '경상북도',
        '경남': '경상남도', '경상남도': '경상남도',
        '대구': '대구광역시', '대구광역시': '대구광역시', '대구시': '대구광역시',
        '울산': '울산광역시', '울산광역시': '울산광역시', '울산시': '울산광역시',
        '부산': '부산광역시', '부산광역시': '부산광역시', '부산시': '부산광역시',
        '제주': '제주특별자치도', '제주도': '제주특별자치도', '제주특별자치도': '제주특별자치도'
    }
    return region_map.get(first, first)


def load_excel_registry(excel_path=EXCEL_PATH):
    """Load and parse preferential enterprise sheets and large corp exclusion from Excel."""
    if not excel_path.exists():
        print(f"[경고] 엑셀 파일이 없습니다: {excel_path}", file=sys.stderr)
        return {}

    import openpyxl  # 빌드에만 필요. 테스트·사이트는 표준 라이브러리만 쓴다.
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    registry = {}  # bizno -> dict

    def clean(v):
        return ' '.join(str(v).split()) if v is not None else ''

    def get_or_create(bizno, name='', region='', ceo='', tel='', address=''):
        if bizno not in registry:
            registry[bizno] = {'bizno': bizno, 'name': name, 'region': region, 'ceo': ceo, 'tel': tel,
                               'address': address, 'excludedAsLargeCorp': False, 'types': [], 'items': []}
        comp = registry[bizno]
        for k, v in (('name', name), ('region', region), ('ceo', ceo), ('tel', tel), ('address', address)):
            if v and not comp[k]:
                comp[k] = v
        return comp

    # 1. 중증('26.06.30)
    if "중증('26.06.30)" in wb.sheetnames:
        ws = wb["중증('26.06.30)"]
        for r in ws.iter_rows(min_row=2, values_only=True):
            bizno = normalize_bizno(r[0])
            if not bizno:
                continue
            name = str(r[1]).strip() if r[1] else ''
            region = extract_region(r[4])
            raw_items = [p.strip() for p in str(r[7]).split(',') if p.strip()] if r[7] else []
            v_from = normalize_date(r[8])
            v_until = normalize_date(r[9])
            status = 'expired' if (v_until and v_until < CHECK_DATE) else 'valid'
            comp = get_or_create(bizno, name, region, ceo=clean(r[2]), tel=clean(r[6]), address=clean(r[5]))
            for it in raw_items:
                if it not in comp['items']:
                    comp['items'].append(it)
            comp['types'].append({
                'type': '중증장애인생산품 생산시설',
                'validFrom': v_from,
                'validUntil': v_until,
                'status': status
            })

    # 2. History sheets with updates (여성, 창업, 장애인)
    history_sheets = [
        ("여성기업('26.06.30)", "여성기업", True),
        ("창업('26.06.30)", "창업기업", False),
        ("장애인기업('26.06.30)", "장애인기업", True)
    ]
    for sname, type_label, has_cancel in history_sheets:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        records = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            bizno = normalize_bizno(r[0])
            if not bizno:
                continue
            name = str(r[1]).strip() if r[1] else ''
            v_from = normalize_date(r[5])
            v_until = normalize_date(r[6])
            cancel = normalize_date(r[7]) if has_cancel and len(r) > 7 else None
            if cancel:
                continue
            key = (v_until or '', v_from or '')
            if bizno not in records or key > records[bizno]['key']:
                records[bizno] = {
                    'name': name,
                    'ceo': clean(r[2]),
                    'validFrom': v_from,
                    'validUntil': v_until,
                    'key': key
                }
        for bizno, rec in records.items():
            v_until = rec['validUntil']
            status = 'expired' if (v_until and v_until < CHECK_DATE) else 'valid'
            comp = get_or_create(bizno, rec['name'], ceo=rec['ceo'])
            comp['types'].append({
                'type': type_label,
                'validFrom': rec['validFrom'],
                'validUntil': v_until,
                'status': status
            })

    # 3. 장애인표준사업장
    std_sname = "장애인표준사업장('26.06.30"
    if std_sname in wb.sheetnames:
        ws = wb[std_sname]
        for r in ws.iter_rows(min_row=2, values_only=True):
            bizno = normalize_bizno(r[0])
            if not bizno:
                continue
            name = str(r[1]).strip() if r[1] else ''
            v_from = normalize_date(r[4])
            region = extract_region(r[7]) if len(r) > 7 else ''
            comp = get_or_create(bizno, name, region, ceo=clean(r[2]), address=clean(r[7]) if len(r) > 7 else '')
            comp['types'].append({
                'type': '장애인표준사업장',
                'validFrom': v_from,
                'validUntil': None,
                'status': 'valid'
            })

    # 4. Simple sheets (사회적, 협동조합, 용사촌, 시범구매)
    simple_sheets = [
        ("사회적('26.06.30)", "사회적기업", 0, 1),
        ("협동조합('26.06.30)", "사회적협동조합", 0, 1),
        ("용사촌('26.06.30)", "자활용사촌", 0, 1),
        ("시범구매(`26.06.30)", "기술개발제품 시범구매", 0, 1)
    ]
    for sname, type_label, col_biz, col_name in simple_sheets:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        seen = set()
        for r in ws.iter_rows(min_row=2, values_only=True):
            bizno = normalize_bizno(r[col_biz])
            if not bizno or bizno in seen:
                continue
            seen.add(bizno)
            name = str(r[col_name]).strip() if len(r) > col_name and r[col_name] else ''
            comp = get_or_create(bizno, name)
            comp['types'].append({
                'type': type_label,
                'validFrom': None,
                'validUntil': None,
                'status': 'valid'
            })

    # 5. 대기업(26.01.12.)+상호출자
    large_sname = "대기업(26.01.12.)+상호출자"
    if large_sname in wb.sheetnames:
        ws = wb[large_sname]
        for r in ws.iter_rows(min_row=2, values_only=True):
            bizno = normalize_bizno(r[1])
            if bizno and bizno in registry:
                registry[bizno]['excludedAsLargeCorp'] = True

    return registry



def load_raw(raw_dir=RAW_DIR):
    """품명별 CSV를 읽어 (품명, 행) 목록으로 돌려준다. shopngCntrctNo+Sno 로 중복 제거."""
    rows, seen = [], set()
    for path in sorted(raw_dir.glob('*.csv')):
        with open(path, encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f):
                key = (r.get('shopngCntrctNo'), r.get('shopngCntrctSno'))
                if key in seen:
                    continue
                seen.add(key)
                rows.append((path.stem, r))
    return rows


def to_iso(s):
    s = (s or '').strip()
    return f'{s[:4]}-{s[4:6]}-{s[6:8]}' if len(s) >= 8 and s[:8].isdigit() else ''


def build_products(raw_rows, registry):
    products = []
    totals = collections.defaultdict(lambda: {'all': 0, 'preferred': 0})
    for query, r in raw_rows:
        totals[query]['all'] += 1
        bizno = normalize_bizno(r.get('cntrctCorpBizno'))
        comp = registry.get(bizno)
        if not comp:
            continue
        totals[query]['preferred'] += 1
        end = to_iso(r.get('cntrctEndDate'))
        products.append({
            'n': r.get('prdctClsfcNoNm', ''),
            'd': r.get('dtilPrdctClsfcNo', ''),
            'id': r.get('prdctIdntNo', ''),
            's': r.get('prdctSpecNm', ''),
            'm': r.get('prdctMakrNm', ''),
            'u': r.get('prdctUnit', ''),
            'p': r.get('cntrctPrceAmt', ''),
            'l': r.get('prdctLrgclsfcNm', ''),
            'b': bizno,
            'cn': r.get('cntrctCorpNm', ''),
            'x': r.get('exclncPrcrmntPrdctYn') == 'Y',
            'k': r.get('smetprCmptProdctYn') == 'Y',
            'e': end,
            'q': query,
        })
    return products, totals


CONTACT_PATH = ROOT / 'data' / 'raw' / 'suppliers.csv'
CONTACT_FIELDS = ('ceo', 'tel', 'fax', 'address', 'zip', 'homepage', 'employees', 'manufacturer', 'opened')


def load_contacts(path=CONTACT_PATH):
    """scripts/collect_suppliers.py 가 만든 업체 기본정보(사용자정보 API). 없으면 빈 dict."""
    if not path.exists():
        return {}
    with open(path, encoding='utf-8-sig', newline='') as f:
        return {normalize_bizno(r['bizno']): r for r in csv.DictReader(f)}


def contact_of(comp, api):
    """엑셀 값과 API 기본정보를 합친다. API 값이 최신이므로 우선하고, 없으면 엑셀 값."""
    a = api or {}
    return {'ceo': a.get('ceo') or comp.get('ceo', ''),
            'tel': a.get('tel') or comp.get('tel', ''),
            'fax': a.get('fax', ''),
            'address': a.get('address') or comp.get('address', ''),
            'zip': a.get('zip', ''),
            'homepage': a.get('homepage', ''),
            'employees': a.get('employees', ''),
            'manufacturer': a.get('manufacturer', ''),
            'opened': a.get('opened', ''),
            'contactSource': 'api' if a else ('excel' if comp.get('tel') or comp.get('address') or comp.get('ceo') else '')}


def build_facilities(registry, contacts=None):
    """종합쇼핑몰에 없는 품목을 위해 중증 시트의 생산품목 목록을 별도로 낸다."""
    out = []
    for comp in registry.values():
        if comp['items']:
            t = next(x for x in comp['types'] if x['type'] == '중증장애인생산품 생산시설')
            out.append({'bizno': comp['bizno'], 'name': comp['name'], 'region': comp['region'],
                        **contact_of(comp, (contacts or {}).get(comp['bizno'])),
                        'items': comp['items'], 'validUntil': t['validUntil'], 'status': t['status'],
                        'otherTypes': sorted({x['type'] for x in comp['types']} - {'중증장애인생산품 생산시설'})})
    return sorted(out, key=lambda x: x['bizno'])


def save(products, facilities, totals, registry, contacts):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    chunk_dir = OUTPUT_DIR / 'chunks'
    chunk_dir.mkdir(exist_ok=True)
    for old in chunk_dir.glob('*.json'):
        old.unlink()
    now = dt.datetime.now(dt.timezone.utc).isoformat()
    used = {p['b'] for p in products}
    suppliers = {b: {'name': registry[b]['name'], 'region': registry[b]['region'],
                     **contact_of(registry[b], contacts.get(b)),
                     'types': registry[b]['types'], 'excludedAsLargeCorp': registry[b]['excludedAsLargeCorp']}
                 for b in sorted(used)}
    by_query = collections.defaultdict(list)
    for p in products:
        by_query[p['q']].append(p)
    chunks = []
    for q, items in sorted(by_query.items()):
        name = f'chunks/{q}.json'
        slim = [{k: v for k, v in it.items() if k != 'q'} for it in items]
        (OUTPUT_DIR / name).write_text(json.dumps({'query': q, 'products': slim}, ensure_ascii=False, separators=(',', ':')) + '\n', encoding='utf-8')
        tc = collections.Counter()
        for b in {it['b'] for it in items}:
            for t in {x['type'] for x in registry[b]['types']}:
                tc[t] += 1
        chunks.append({'query': q, 'file': name, 'count': len(items), 'all': totals[q]['all'],
                       'suppliers': len({it['b'] for it in items}), 'typeSuppliers': dict(tc.most_common()),
                       'names': sorted({it['n'] for it in items})})
    catalog = {'schemaVersion': 2, 'basisDate': '2026-06-30', 'checkDate': CHECK_DATE, 'generatedAt': now,
               'source': 'https://www.data.go.kr/data/15129471/openapi.do', 'totalCount': len(products),
               'fields': {'n': '품명', 'd': '세부품명번호', 'id': '물품식별번호', 's': '규격', 'm': '제조사', 'u': '단위',
                          'p': '계약단가', 'l': '대분류', 'b': '사업자등록번호', 'cn': '계약업체명',
                          'x': '우수제품', 'k': '중기간경쟁제품', 'e': '계약종료일'},
               'chunks': chunks, 'suppliers': suppliers, 'facilities': facilities}
    CATALOG_PATH.write_text(json.dumps(catalog, ensure_ascii=False, separators=(',', ':')) + '\n', encoding='utf-8')
    types = collections.Counter()
    corps = collections.defaultdict(set)
    for p in products:
        for t in {x['type'] for x in registry[p['b']]['types']}:
            types[t] += 1
            corps[t].add(p['b'])
    summary = {'schemaVersion': 2, 'basisDate': '2026-06-30', 'checkDate': CHECK_DATE, 'generatedAt': now,
               'totalProducts': len(products), 'uniqueSuppliers': len(used), 'facilities': len(facilities),
               'typeDistribution': {t: {'products': c, 'suppliers': len(corps[t])} for t, c in types.most_common()},
               'queryTotals': dict(sorted(totals.items())),
               'expiredCount': sum(1 for p in products if p['e'] and p['e'] < CHECK_DATE),
               'suppliersWithContact': sum(1 for s in suppliers.values() if s['tel'] or s['address']),
               'catalogBytes': CATALOG_PATH.stat().st_size,
               'chunkBytes': sum((OUTPUT_DIR / c['file']).stat().st_size for c in chunks)}
    SUMMARY_PATH.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return summary


def main():
    registry = load_excel_registry(EXCEL_PATH)
    print(f'엑셀 우대기업 {len(registry):,}곳')
    raw_rows = load_raw(RAW_DIR)
    if not raw_rows:
        raise ValueError('data/raw/shopmall/에 수집된 CSV가 없습니다. scripts/collect_shopmall.py를 먼저 실행하세요. '
                         '실제 업체명·사업자등록번호에 가공한 물품·가격을 붙여 게시하지 않습니다.')
    print(f'종합쇼핑몰 원시 {len(raw_rows):,}건')
    products, totals = build_products(raw_rows, registry)
    contacts = load_contacts()
    print(f'업체 기본정보(연락처) {len(contacts):,}곳')
    facilities = build_facilities(registry, contacts)
    summary = save(products, facilities, totals, registry, contacts)
    print(f"우대기업 품목 {summary['totalProducts']:,}건 / 업체 {summary['uniqueSuppliers']:,}곳 / 중증 생산시설 {summary['facilities']:,}곳")
    print(f"catalog.json {summary['catalogBytes']/1024:.0f}KB / 청크 합계 {summary['chunkBytes']/1048576:.1f}MB (품명 선택 시 개별 로드)")
    for t, v in summary['typeDistribution'].items():
        print(f"   {t:14s} {v['products']:7,}건 {v['suppliers']:5,}곳")


if __name__ == '__main__':
    try:
        main()
    except (ValueError, OSError) as error:
        print(f'중단: {error}', file=sys.stderr)
        sys.exit(1)
